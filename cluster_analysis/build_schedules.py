#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""يبني صفحة «الجداول» (schedules.html):
- ملخص البطولات + قوالب الجداول حسب عدد الفرق (من ملف الجداول)
- مجموعات تحت 5 الفعلية بأسماء فرقها؛ ومن حجمه قالب يُعرض جدوله بالأسماء.
المدخلات: schedules_u5.xlsx (القوالب) + schedule_groups_u5.json (المجموعات→الفرق)."""
import json, os, re
import openpyxl

BASE = os.path.dirname(os.path.abspath(__file__)) + "/"
XLSX = BASE + "schedules_u5.xlsx"
GROUPS = BASE + "schedule_groups_u5.json"

wb = openpyxl.load_workbook(XLSX, read_only=True, data_only=True)

def S(c):
    return "" if c is None else str(c).strip()

# ---- الإعدادات (المدخلات) ----
settings = []
if "المدخلات" in wb.sheetnames:
    for r in wb["المدخلات"].iter_rows(values_only=True):
        k, v = S(r[0]), (S(r[1]) if len(r) > 1 else "")
        if k and v and k not in ("المدخلات", "البند"):
            settings.append([k, v])

# ---- ملخص البطولات ----
summary = {"headers": [], "rows": []}
if "ملخص البطولات" in wb.sheetnames:
    ws = wb["ملخص البطولات"]
    rows = [[S(c) for c in r] for r in ws.iter_rows(values_only=True)]
    hi = next((i for i, r in enumerate(rows) if r and r[0] == "عدد الفرق"), None)
    if hi is not None:
        summary["headers"] = [c for c in rows[hi] if c]
        ncol = len(summary["headers"])
        for r in rows[hi + 1:]:
            if not r or not r[0]:
                break
            if not re.match(r"^\d", r[0]):
                break
            summary["rows"].append(r[:ncol])

# ---- قوالب الجداول حسب عدد الفرق ----
def parse_template(ws):
    rows = [[S(c) for c in r] for r in ws.iter_rows(values_only=True)]
    title = rows[0][0] if rows and rows[0] else ""
    hi = next((i for i, r in enumerate(rows)
               if len(r) > 1 and r[0] == "اليوم" and r[1].startswith("الوقت")), None)
    if hi is None:
        return None
    hdr = rows[hi]
    pitch_idx = [j for j, c in enumerate(hdr) if c.startswith("ملعب")]
    pitches = [hdr[j] for j in pitch_idx]
    days, cur, slots = [], None, None
    for r in rows[hi + 1:]:
        day, time = (r[0] if r else ""), (r[1] if len(r) > 1 else "")
        if day.startswith("اليوم"):
            if cur is not None:
                days.append({"day": cur, "slots": slots})
            cur, slots = day, []
        # صف فترة حقيقي: الوقت يحوي ":" مثل «5:00–5:15 م» (يستبعد جداول الإحصاءات الملحقة)
        if time and ":" in time and cur is not None:
            cells = ["" if S(r[j]) == "0" else (r[j] if j < len(r) else "")
                     for j in pitch_idx]
            slots.append({"time": time, "cells": cells})
    if cur is not None:
        days.append({"day": cur, "slots": slots})
    # إزالة صفوف الاستراحة/الفارغة من نهاية كل يوم (لا معنى لاستراحة في الآخر)
    def has_match(cells):
        return any(re.match(r"^\d+\s*ضد\s*\d+$", S(c)) for c in cells)
    for d in days:
        while d["slots"] and not has_match(d["slots"][-1]["cells"]):
            d["slots"].pop()
    days = [d for d in days if d["slots"]]
    return {"title": title, "pitches": pitches, "days": days}

templates = {}
for sn in wb.sheetnames:
    m = re.match(r"^(\d+)\s*(?:فرق|فريقًا)(?:\s*\(\d+\))?$", sn)
    if not m:
        continue
    size = int(m.group(1))
    if size in templates:        # تجاهل النسخ المكررة مثل «6 فرق (2)»
        continue
    t = parse_template(wb[sn])
    if t:
        templates[size] = t

# ---- المجموعات الفعلية ----
G = json.load(open(GROUPS, encoding="utf-8"))
avail = set(templates.keys())
groups = []
for g in G["groups"]:
    teams = [t["name"] for t in g["teams"]]
    size = len(teams)
    groups.append({"group": g["group"], "region": g.get("region", ""),
                   "teams": teams, "size": size,
                   "tmpl": size if size in avail else None})

DATA = {"age": G.get("age", "تحت 5"), "settings": settings, "summary": summary,
        "templates": {str(k): v for k, v in templates.items()}, "groups": groups}

HTML = r"""<!DOCTYPE html>
<html lang="ar" dir="rtl"><head>
<meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>جداول البطولات — __AGE__</title>
<link href="https://fonts.googleapis.com/css2?family=Tajawal:wght@400;500;700;800&display=swap" rel="stylesheet">
<style>
*{box-sizing:border-box} body{margin:0;font-family:'Tajawal',sans-serif;background:#04150e;color:#eafff3}
.top{background:#006C35;padding:14px 20px;display:flex;flex-wrap:wrap;gap:12px;align-items:center}
.top img.logo{height:46px}
.top h1{font-size:19px;margin:0;font-weight:800}
.wrap{max-width:1080px;margin:0 auto;padding:16px}
.cards{display:flex;flex-wrap:wrap;gap:10px;margin-bottom:14px}
.card{background:#0d3f2d;border:1px solid #1c6b49;border-radius:12px;padding:12px 16px;flex:1;min-width:150px}
.card .l{color:#8fdcb4;font-size:12px} .card .v{font-size:18px;font-weight:800;margin-top:3px}
.tabs{display:flex;gap:8px;flex-wrap:wrap;margin-bottom:14px}
.tab{background:#0d4b32;border:1px solid #1c7a52;color:#eafff3;border-radius:20px;padding:8px 18px;cursor:pointer;font-family:'Tajawal';font-size:14px;font-weight:700}
.tab.on{background:#ffd166;color:#04150e;border-color:#ffd166}
.panel{background:#0b2c1f;border:1px solid #14543a;border-radius:14px;padding:14px}
.sel{background:#0d3f2d;color:#eafff3;border:1px solid #1c6b49;border-radius:10px;padding:10px 14px;font-family:'Tajawal';font-weight:700;font-size:14px;min-width:min(300px,100%);max-width:100%;margin-bottom:12px;cursor:pointer}
.sel option{background:#0b2c1f;color:#eafff3}
.glist{display:flex;flex-wrap:wrap;gap:8px}
.gbtn{background:#0d3f2d;border:1px solid #1c6b49;border-radius:10px;padding:9px 13px;cursor:pointer;color:#eafff3;font-family:'Tajawal';font-weight:700;font-size:13px;text-align:right}
.gbtn:hover{background:#11523a} .gbtn.on{background:#ffd166;color:#04150e;border-color:#ffd166}
.gbtn .c{font-size:11px;font-weight:500;color:#9fdca0;display:block;margin-top:2px}
.gbtn.on .c{color:#0a3b23}
.roster{display:flex;flex-wrap:wrap;gap:6px;margin:10px 0 14px}
.rteam{background:#0d3f2d;border:1px solid #1c6b49;border-radius:8px;padding:5px 10px;font-size:12.5px}
.rteam b{color:#ffd166;margin-left:5px}
.rteam.pick{cursor:pointer} .rteam.pick:hover{background:#11523a}
.rteam.pick.on{background:#ffd166;color:#04150e;border-color:#ffd166} .rteam.pick.on b{color:#0a3b23}
h3.sec{color:#ffd166;font-size:15px;margin:16px 0 8px;border-bottom:1px solid #14543a;padding-bottom:5px}
.dayttl{color:#8fdcb4;font-weight:800;font-size:14px;margin:14px 0 6px}
.subttl{color:#ffd166;font-weight:800;font-size:14px;margin:12px 0 4px}
table{border-collapse:collapse;width:100%;font-size:12.5px;margin-bottom:6px}
th,td{border:1px solid #14543a;padding:6px 8px;text-align:center}
th{background:#0d4b32;color:#eafff3;font-weight:700} td{background:#0a2418}
td.time{background:#0d3f2d;color:#ffd166;font-weight:700;white-space:nowrap}
td.rest{color:#6f9a86;font-style:italic;background:#08190f}
.mtch b{color:#bfe9d4}
.note{color:#ffcf8a;background:#3a2f00;border:1px solid #ffd16655;border-radius:8px;padding:8px 11px;font-size:12.5px;margin-top:8px}
.muted{color:#8fdcb4;font-size:13px}
.stbl{overflow-x:auto}
.hint{color:#8fdcb4;font-size:12.5px;margin:2px 0 12px}
</style></head><body>
<div class="top">
 <img class="logo" src="logo.png" alt="الاتحاد" onerror="this.remove()">
 <h1>🗓️ جداول البطولات — __AGE__</h1>
</div>
<div class="wrap">
 <div class="cards" id="cards"></div>
 <div class="tabs" id="tabs"></div>
 <div class="panel" id="panel"></div>
</div>
<script>
const D=__DATA__;
function gp(n){return String(n).replace(/\B(?=(\d{3})+(?!\d))/g,',');}
function arCount(n,one,two,few,many){const m=n%100;if(n===1)return one;if(n===2)return two;
  if(m>=3&&m<=10)return gp(n)+' '+few;return gp(n)+' '+many;}
function nTeam(n){return arCount(n,'فريق واحد','فريقان','فرق','فريقًا');}
function nMatch(n){return arCount(n,'مباراة واحدة','مباراتان','مباريات','مباراةً');}
function esc(s){return String(s).replace(/[&<>]/g,c=>({'&':'&amp;','<':'&lt;','>':'&gt;'}[c]));}
function dayNum(d){return String(d).replace(/^\s*اليوم\s*/,'');}
function parseSubgroups(title){const subs=[];['أ','ب','ج','د'].forEach(l=>{
  const m=title.match(new RegExp(l+'\\s*:\\s*(\\d+)\\s*[-\\u2013]\\s*(\\d+)'));
  if(m)subs.push({label:l,from:+m[1],to:+m[2]});});return subs;}
function subLabelOf(subs,num){for(const s of subs){if(num>=s.from&&num<=s.to)return s.label;}return '';}
// استبدال «س ضد ص» بأسماء الفرق
function cellHTML(cell,teams){
  const m=cell.match(/^(\d+)\s*ضد\s*(\d+)$/);
  if(m&&teams){const a=teams[+m[1]-1]||('#'+m[1]), b=teams[+m[2]-1]||('#'+m[2]);
    return '<span class="mtch"><b>'+esc(a)+'</b> ضد <b>'+esc(b)+'</b></span>';}
  if(/^استراحة/.test(cell))return '<span class="rest-in">—</span>';
  return esc(cell);
}
function isRestRow(s){return s.cells.some(c=>/^استراحة/.test(c))&&s.cells.every(c=>c===''||/^استراحة/.test(c));}
function renderTemplate(t,teams){
  if(!t)return '';
  let h='';
  t.days.forEach(day=>{
    h+='<div class="dayttl">'+esc(day.day)+'</div><div class="stbl"><table><tr><th>الوقت</th>'+
      t.pitches.map(p=>'<th>'+esc(p)+'</th>').join('')+'</tr>';
    day.slots.forEach(s=>{
      if(isRestRow(s)){h+='<tr><td class="time">'+esc(s.time)+'</td><td class="rest" colspan="'+t.pitches.length+'">استراحة</td></tr>';return;}
      h+='<tr><td class="time">'+esc(s.time)+'</td>'+
        s.cells.map(c=>'<td>'+(c?cellHTML(c,teams):'')+'</td>').join('')+'</tr>';
    });
    h+='</table></div>';
  });
  return h;
}
// بطاقات الإعدادات
document.getElementById('cards').innerHTML=D.settings.map(s=>
  '<div class="card"><div class="l">'+esc(s[0])+'</div><div class="v">'+esc(s[1])+'</div></div>').join('');
// التبويبات
let tab='groups';
const TABS=[['groups','المجموعات'],['tmpl','قوالب الجداول'],['summary','ملخص البطولات']];
let curG=D.groups.length?0:-1, curT=Object.keys(D.templates).map(Number).sort((a,b)=>a-b)[0];
let curTeam=null; // فهرس الفريق المختار داخل المجموعة (null = الجدول الكامل)
// مباريات فريق واحد داخل قالب المجموعة
function teamMatches(t,teams,idx){
  const num=idx+1, out=[];
  t.days.forEach(day=>day.slots.forEach(s=>s.cells.forEach((c,ci)=>{
    const m=c&&c.match(/^(\d+)\s*ضد\s*(\d+)$/); if(!m)return;
    const a=+m[1], b=+m[2];
    if(a===num||b===num){const opp=a===num?b:a;
      out.push({day:day.day,time:s.time,pitch:t.pitches[ci],opp:teams[opp-1]||('#'+opp)});}
  })));
  return out;
}
function teamMatchesHTML(t,teams,idx){
  const rows=teamMatches(t,teams,idx);
  let h='<div class="stbl"><table><tr><th>اليوم</th><th>الوقت</th><th>الملعب</th><th>الخصم</th></tr>';
  rows.forEach(r=>{h+='<tr><td>'+esc(dayNum(r.day))+'</td><td class="time">'+esc(r.time)+'</td><td>'+esc(r.pitch)+'</td><td class="mtch"><b>'+esc(r.opp)+'</b></td></tr>';});
  h+='</table></div>';
  return {html:h,count:rows.length};
}
function renderTabs(){document.getElementById('tabs').innerHTML=TABS.map(([k,l])=>
  '<button class="tab'+(k===tab?' on':'')+'" data-k="'+k+'">'+l+'</button>').join('');
  document.querySelectorAll('#tabs .tab').forEach(b=>b.onclick=()=>{tab=b.dataset.k;render();});}
function groupsPanel(){
  const g=D.groups[curG];
  let h='<div class="hint">💡 اختر مجموعة لعرض فرقها وجدولها. الفرق المرقّمة مؤقتة (ترتيب التسجيل).</div>';
  h+='<select id="gsel" class="sel">'+D.groups.map((x,i)=>'<option value="'+i+'"'+(i===curG?' selected':'')+'>'+
    esc(x.group)+' — '+nTeam(x.size)+'</option>').join('')+'</select>';
  if(g){
    h+='<h3 class="sec">'+esc(g.group)+' — '+nTeam(g.size)+(g.region?' · '+esc(g.region):'')+'</h3>';
    if(g.tmpl){
      const tmpl=D.templates[String(g.tmpl)];
      const subs=parseSubgroups(tmpl.title);
      if(subs.length>1){
        h+='<div class="hint">هذه المجموعة مقسّمة إلى '+subs.length+' مجموعات فرعية، تلعب كل واحدة على ملاعبها:</div>';
        subs.forEach(sg=>{h+='<div class="subttl">المجموعة '+sg.label+'</div><div class="roster">'+
          g.teams.slice(sg.from-1,sg.to).map((t,i)=>'<span class="rteam"><b>'+(sg.from+i)+'</b>'+esc(t)+'</span>').join('')+'</div>';});
      }
      h+='<div class="hint">اختر فريقًا لعرض مبارياته فقط، أو «الجدول الكامل» للكل:</div>';
      h+='<select id="tsel" class="sel"><option value="-1"'+(curTeam===null?' selected':'')+'>📋 الجدول الكامل</option>'+
        g.teams.map((t,i)=>{const sl=subLabelOf(subs,i+1);return '<option value="'+i+'"'+(curTeam===i?' selected':'')+'>'+(i+1)+' — '+esc(t)+(sl?' (المجموعة '+sl+')':'')+'</option>';}).join('')+'</select>';
      if(curTeam!==null){
        const r=teamMatchesHTML(tmpl,g.teams,curTeam);const sl=subLabelOf(subs,curTeam+1);
        h+='<h3 class="sec">مباريات '+esc(g.teams[curTeam])+(sl?' (المجموعة '+sl+')':'')+' — '+nMatch(r.count)+'</h3>'+r.html;
      }else{
        h+=renderTemplate(tmpl,g.teams);
      }
    }else{
      h+='<div class="roster">'+g.teams.map((t,i)=>'<span class="rteam"><b>'+(i+1)+'</b>'+esc(t)+'</span>').join('')+'</div>';
      h+='<div class="note">لا يوجد قالب جدول جاهز لعدد '+nTeam(g.size)+' في ملف الجداول بعد — عُرضت قائمة الفرق فقط.</div>';
    }
  }
  return h;
}
function tmplPanel(){
  const sizes=Object.keys(D.templates).map(Number).sort((a,b)=>a-b);
  let h='<div class="hint">💡 قوالب الجداول حسب عدد الفرق (بالأرقام). تُطبَّق على أي مجموعة بنفس العدد.</div>';
  h+='<select id="tmsel" class="sel">'+sizes.map(s=>'<option value="'+s+'"'+(s===curT?' selected':'')+'>'+nTeam(s)+'</option>').join('')+'</select>';
  const t=D.templates[String(curT)];
  if(t){const subs=parseSubgroups(t.title);
    h+='<h3 class="sec">'+esc(t.title)+'</h3>';
    if(subs.length>1)h+='<div class="hint">أرقام '+subs.map(s=>'المجموعة '+s.label+': '+s.from+'–'+s.to).join(' · ')+'</div>';
    h+=renderTemplate(t,null);}
  return h;
}
function summaryPanel(){
  const s=D.summary;if(!s.headers.length)return '<div class="muted">لا يوجد ملخص.</div>';
  let h='<h3 class="sec">ملخص جداول المجموعات</h3><div class="stbl"><table><tr>'+
    s.headers.map(x=>'<th>'+esc(x)+'</th>').join('')+'</tr>';
  s.rows.forEach(r=>{h+='<tr>'+s.headers.map((_,j)=>'<td'+(j===0?' class="time"':'')+'>'+esc(r[j]||'')+'</td>').join('')+'</tr>';});
  h+='</table></div>';return h;
}
function render(){
  renderTabs();
  const p=document.getElementById('panel');
  p.innerHTML = tab==='groups'?groupsPanel() : tab==='tmpl'?tmplPanel() : summaryPanel();
  if(tab==='groups'){
    const gs=document.getElementById('gsel');if(gs)gs.onchange=()=>{curG=+gs.value;curTeam=null;render();window.scrollTo(0,0);};
    const ts=document.getElementById('tsel');if(ts)ts.onchange=()=>{const v=+ts.value;curTeam=(v<0?null:v);render();};
  }
  if(tab==='tmpl'){const ms=document.getElementById('tmsel');if(ms)ms.onchange=()=>{curT=+ms.value;render();window.scrollTo(0,0);};}
}
render();
</script>
</body></html>"""

HTML = HTML.replace("__AGE__", DATA["age"]).replace("__DATA__", json.dumps(DATA, ensure_ascii=False))
open(BASE + "schedules.html", "w", encoding="utf-8").write(HTML)
print("saved schedules.html", len(HTML), "bytes | قوالب:", sorted(templates.keys()),
      "| مجموعات:", len(groups), "لها قالب:", sum(1 for g in groups if g["tmpl"]))
