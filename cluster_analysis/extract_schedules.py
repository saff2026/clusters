#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""يستخرج بيانات القوالب من ملف جداول (يدعم الصيغتين: أعمدة ملاعب، أو صف لكل مباراة).
الاستخدام: python3 extract_schedules.py <ملف.xlsx> <out.json>"""
import json, os, re, sys
import openpyxl

XLSX, OUT = sys.argv[1], sys.argv[2]
wb = openpyxl.load_workbook(XLSX, read_only=True, data_only=True)

def S(c):
    return "" if c is None else str(c).strip()

# ---- الإعدادات ----
settings = []
if "المدخلات" in wb.sheetnames:
    for r in wb["المدخلات"].iter_rows(values_only=True):
        k = S(r[0]); v = S(r[1]) if len(r) > 1 else ""
        if k and v and "المدخلات" not in k and k != "البند":
            settings.append([k, v])

# ---- الملخص ----
summary = {"headers": [], "rows": []}
sumsheet = next((s for s in ("الملخص", "ملخص البطولات") if s in wb.sheetnames), None)
if sumsheet:
    rows = [[S(c) for c in r] for r in wb[sumsheet].iter_rows(values_only=True)]
    hi = next((i for i, r in enumerate(rows) if r and r[0] == "عدد الفرق"), None)
    if hi is not None:
        summary["headers"] = [c for c in rows[hi] if c]
        n = len(summary["headers"])
        for r in rows[hi + 1:]:
            if not r or not r[0] or not re.match(r"^\d", r[0]):
                break
            summary["rows"].append(r[:n])

def subgroups_from_matches(days):
    """المجموعات الفرعية = مكوّنات مترابطة من مباريات القالب (من يلعب مع من)."""
    adj = {}
    for d in days:
        for s in d["slots"]:
            for c in s["cells"]:
                m = re.match(r"^(\d+)\s*ضد\s*(\d+)$", c)
                if m:
                    a, b = int(m.group(1)), int(m.group(2))
                    adj.setdefault(a, set()).add(b)
                    adj.setdefault(b, set()).add(a)
    seen, comps = set(), []
    for node in sorted(adj):
        if node in seen:
            continue
        stack, comp = [node], []
        while stack:
            x = stack.pop()
            if x in seen:
                continue
            seen.add(x); comp.append(x)
            stack.extend(adj[x] - seen)
        comps.append(sorted(comp))
    comps.sort(key=lambda c: c[0])
    labels = ["أ", "ب", "ج", "د", "هـ", "و"]
    return [{"label": labels[i] if i < len(labels) else str(i + 1), "teams": comp}
            for i, comp in enumerate(comps)]

def parse_wide(rows, hi):
    hdr = rows[hi]
    pidx = [j for j, c in enumerate(hdr) if c.startswith("ملعب")]
    pitches = [hdr[j] for j in pidx]
    days, cur, slots = [], None, None
    for r in rows[hi + 1:]:
        day = r[0] if r else ""; time = r[1] if len(r) > 1 else ""
        if day.startswith("اليوم"):
            if cur is not None: days.append({"day": cur, "slots": slots})
            cur, slots = day, []
        if time and ":" in time and cur is not None:
            cells = ["" if S(r[j]) == "0" else (r[j] if j < len(r) else "") for j in pidx]
            slots.append({"time": time, "cells": cells})
    if cur is not None: days.append({"day": cur, "slots": slots})
    return pitches, days

def parse_long(rows, hi):
    """صف لكل مباراة: اليوم | الوقت | الملعب | المباراة."""
    hdr = rows[hi]
    ci = {name: j for j, c in enumerate(hdr) for name in ["اليوم", "الوقت", "الملعب", "المباراة"] if c == name}
    jd, jt, jp, jm = ci["اليوم"], ci["الوقت"], ci["الملعب"], ci["المباراة"]
    # اجمع الملاعب الفريدة بالترتيب الرقمي
    pnums = []
    recs = []
    for r in rows[hi + 1:]:
        day, time, pit, mt = S(r[jd]), S(r[jt]), S(r[jp]), S(r[jm])
        if not (day or time): continue
        if not time or ":" not in time: continue
        recs.append((day, time, pit, mt))
        if pit and pit not in pnums: pnums.append(pit)
    def pk(p):
        m = re.search(r"\d+", p); return int(m.group()) if m else 0
    pnums.sort(key=pk)
    pitches = ["ملعب " + p if not p.startswith("ملعب") else p for p in pnums]
    pindex = {p: i for i, p in enumerate(pnums)}
    # جمّع حسب (يوم، وقت)
    days = []
    order = []
    dmap = {}
    for day, time, pit, mt in recs:
        d = dmap.get(day)
        if d is None:
            d = {"day": day, "_slots": {}, "_order": []}; dmap[day] = d; order.append(day)
        if time not in d["_slots"]:
            d["_slots"][time] = [""] * len(pnums); d["_order"].append(time)
        if pit in pindex:
            d["_slots"][time][pindex[pit]] = mt
    for day in order:
        d = dmap[day]
        slots = [{"time": t, "cells": d["_slots"][t]} for t in d["_order"]]
        days.append({"day": day, "slots": slots})
    return pitches, days

def parse_template(ws):
    rows = [[S(c) for c in r] for r in ws.iter_rows(values_only=True)]
    title = rows[0][0] if rows and rows[0] else ""
    # نوع الصيغة من صف الترويسة
    hi = next((i for i, r in enumerate(rows)
               if len(r) > 1 and r[0] == "اليوم" and r[1].startswith("الوقت")), None)
    if hi is None:
        return None
    hdr = rows[hi]
    is_long = any(c == "المباراة" for c in hdr) and not any(c.startswith("ملعب") for c in hdr)
    pitches, days = (parse_long if is_long else parse_wide)(rows, hi)
    # قصّ صفوف الاستراحة/الفارغة من نهاية كل يوم
    def has_match(cells): return any(re.match(r"^\d+\s*ضد\s*\d+$", S(c)) for c in cells)
    for d in days:
        while d["slots"] and not has_match(d["slots"][-1]["cells"]): d["slots"].pop()
    days = [d for d in days if d["slots"]]
    return {"title": title, "pitches": pitches, "subgroups": subgroups_from_matches(days), "days": days}

templates = {}
for sn in wb.sheetnames:
    m = re.match(r"^(\d+)\s*(?:فرق|فريقًا|فريق)(?:\s*\(\d+\))?$", sn)
    if not m:
        continue
    size = int(m.group(1))
    if size in templates:
        continue
    t = parse_template(wb[sn])
    if t and t["days"]:
        templates[size] = t

src = {"settings": settings, "summary": summary,
       "templates": {str(k): v for k, v in templates.items()}}
json.dump(src, open(OUT, "w", encoding="utf-8"), ensure_ascii=False)
print("saved", os.path.basename(OUT), "| قوالب:", sorted(templates.keys()),
      "| متعددة المجموعات:", [k for k, v in templates.items() if len(v["subgroups"]) > 1])
